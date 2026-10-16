"""Build an editable report from the supplied college template and a slide deck."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from pptx import Presentation
from pptx.util import Inches as PptInches, Pt as PptPt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

ROOT=Path(__file__).resolve().parents[1]
TEMPLATE=ROOT/'docs'/'PBL Report Template.docx'
OUT=ROOT/'docs'/'final_report.docx'
PPT=ROOT/'presentation'/'payroll_pbl_presentation.pptx'
TITLE='Design and Implementation of a Database Management System for Employee Payroll and Statutory Deduction Management System'

def build_report():
    doc=Document(TEMPLATE)
    body=doc._element.body
    for child in list(body):
        if child.tag.endswith('}sectPr'): continue
        body.remove(child)
    sec=doc.sections[0]
    sec.left_margin=Inches(1.25);sec.right_margin=Inches(1);sec.top_margin=Inches(1);sec.bottom_margin=Inches(1)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12)
    normal.paragraph_format.line_spacing=1.5
    normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    for nm in ('Heading 1','Heading 2','Heading 3'):
        st=doc.styles[nm];st.font.name='Times New Roman';st.font.bold=True;st.font.size=Pt(14 if nm=='Heading 1' else 12)

    def p(text='',bold=False,center=False,italic=False):
        q=doc.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
        r=q.add_run(text);r.bold=bold;r.italic=italic;r.font.name='Times New Roman';r.font.size=Pt(12)
        return q
    def h(text,level=1):
        q=doc.add_paragraph(style=f'Heading {level}');q.paragraph_format.keep_with_next=True
        r=q.add_run(text.upper() if level==1 else text);r.font.name='Times New Roman';r.font.size=Pt(14 if level==1 else 12);r.bold=True
        return q
    def bullet(text):
        q=doc.add_paragraph();q.paragraph_format.line_spacing=1.5;q.paragraph_format.left_indent=Inches(.25);q.add_run('• '+text)
    def numbered(text):
        number=getattr(numbered,'counter',0)+1;numbered.counter=number
        q=doc.add_paragraph();q.paragraph_format.line_spacing=1.5;q.paragraph_format.left_indent=Inches(.25);q.add_run(f'{number}. {text}')
    def table(headers,rows):
        t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER
        for i,v in enumerate(headers):t.rows[0].cells[i].text=str(v)
        for row in rows:
            cells=t.add_row().cells
            for i,v in enumerate(row):cells[i].text=str(v)
        for row in t.rows:
            for cell in row.cells:
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for para in cell.paragraphs:
                    para.paragraph_format.line_spacing=1.5
                    for run in para.runs:run.font.name='Times New Roman';run.font.size=Pt(12)
        for cell in t.rows[0].cells:
            for run in cell.paragraphs[0].runs:run.bold=True
        doc.add_paragraph()
        return t
    def page():doc.add_page_break()

    # Cover and required preliminary pages.
    p('Woxsen University',True,True);p('School of Technology',False,True);p('',center=True);p('B.Tech – Computer Science and Engineering',True,True)
    p('PROJECT BASED LEARNING REPORT',True,True);p('TITLE OF THE PROJECT',True,True);p(TITLE,True,True)
    p('Course Code: 25TU03MJM0',center=True);p('Course Name: Database Management Systems',center=True);p('Semester: 3   Academic Year: 2026–27',center=True)
    p('Project Number: [ENTER PROJECT NUMBER]',center=True)
    p('Submitted by:',True,True);p('Name: Siddharth [FULL NAME]',center=True);p('Register Number: [ROLL NUMBER]',center=True);p('Section: CSE Panthers',center=True)
    p('Under the Guidance of:',True,True);p('Faculty Name: Dr. Sahithi',center=True);p('Department of Computer Science and Engineering',center=True);p('2026',center=True)
    page();h('Declaration')
    p(f'I hereby declare that the project titled “{TITLE}” submitted by me is an original work carried out under the guidance of the faculty mentioned above. I confirm that this work has not been submitted in part or full for another academic qualification. Sources used are acknowledged. I will complete any institution-required originality checks before final submission.')
    p('Place: ____________________     Date: ____________________');p('Signature of Candidate: ____________________');p('Name: Siddharth [FULL NAME]');p('Register Number: [ROLL NUMBER]')
    page();h('Certificate')
    p(f'This is to certify that the project titled “{TITLE}” is a bonafide work carried out by Siddharth [FULL NAME] (Register No: [ROLL NUMBER]), B.Tech Computer Science and Engineering, during Academic Year 2026–27, under the guidance of Dr. Sahithi.')
    p('Guide Signature: ____________________');p('Dr. Sahithi');p('Designation: ____________________');p('Department of Computer Science and Engineering')
    p('Institutional signature and certification are intentionally left for the authorized faculty.')
    page();h('Acknowledgement')
    p('I express my sincere thanks to Dr. Sahithi for guidance and feedback during this database project. I thank the School of Technology and the Department of Computer Science and Engineering for the course structure and project-based learning opportunity. I also acknowledge the authors and documentation teams cited in the references. This report and its software are an educational prototype; the student must verify the database behavior, add genuine screenshots and replace the remaining identity placeholders before submission.')
    page();h('Abstract')
    abstract=('Organizations need reliable records for employees, salary components, attendance, statutory deductions, loans and monthly payments. Spreadsheet-based payroll work can duplicate data, propagate formula errors and make it difficult to explain how net pay was derived. This project designs and implements a relational database prototype for employee payroll and deduction management using MySQL 8.0 and Python. The conceptual model includes departments, designations, employees, effective-dated salary structures, attendance summaries, payroll periods, payroll details, earnings, deductions, loans, payslips and bank-transfer instructions. The relational design separates repeating earning and deduction items and stores organizational descriptions with their own entities. Unique and foreign-key constraints protect identities and relationships; CHECK constraints validate local numeric and date ranges; triggers enforce selected cross-row payroll rules. A Tkinter application supports employee maintenance, search and reports. A Python payroll module calculates an employee-period payroll inside a transaction, records component lines and updates loan balances together. SQL scripts provide fictional sample data and reports using joins, grouping, filtering, nested queries, views and rollback. The prototype rates are illustrative classroom assumptions and are not jurisdiction-approved statutory calculations. The bank transfer is only a database instruction record; no real bank integration is implemented. This workspace did not have a MySQL server, so database execution, screenshots and empirical performance results remain to be collected by the student. The design and code provide a practical foundation for demonstrating normalization, constraints, transaction handling and application integration, while the test matrix states expected outcomes without presenting unrun cases as passed. Deliverables are arranged to support the three scheduled project reviews.')
    p(abstract)
    page();h('Table of Contents')
    for item in ['1. Introduction','2. Problem Statement and Objectives','3. Literature Review and Background','4. Methodology and System Design','5. Implementation','6. Results and Discussion','7. Conclusion and Future Scope','References','Appendix A – Setup and Demonstration'] : p(item)
    page();h('List of Figures');p('Figure 1. System architecture (Chapter 4)');p('Figure 2. Payroll processing flow (Chapter 4)');p('Figure 3. ER diagram (see docs/review-1/er_diagram.md)');p('Figure 4. Application screen captures (to be inserted after local run)')
    page();h('List of Tables');p('Table 1. Related-work comparison');p('Table 2. Functional requirements');p('Table 3. Entity data dictionary');p('Table 4. Normalization summary');p('Table 5. Test matrix');p('Table 6. Objective-result mapping')
    page()

    h('1. Introduction')
    h('1.1 Background',2)
    p('Payroll is a recurring business process that converts employee and work-period information into a payment amount. The calculation may combine basic salary, allowances, overtime, statutory deductions, voluntary deductions and loan repayments. The result must be traceable to its employee, period and component inputs. Even a small organization needs a reliable way to distinguish current employee details from historical payroll facts.')
    p('A relational database stores related facts in tables and links them through keys. Codd’s relational model provides the basis for representing data as relations and querying them declaratively [1]. Normalization organizes relations to reduce update anomalies [2], [3]. These concepts suit payroll because employee, department, salary, attendance and payment information have distinct lifecycles and relationships.')
    h('1.2 Motivation and real-world relevance',2)
    p('A copied spreadsheet formula may omit a row or use a different percentage from the approved rule. Duplicate employee-period entries can produce duplicate payment instructions. A stored net-pay total without its earning and deduction components is hard to audit. Database constraints can prevent many invalid states at the point of storage, while transactions can ensure a group of related writes completes together [4], [5].')
    p('The project applies these ideas in a bounded coursework system. It keeps sample bank accounts fictional, uses illustrative statutory percentages and does not connect to a real bank or government service. This boundary allows the student to focus on database correctness without presenting the example as production payroll policy.')
    h('1.3 Scope and organization',2)
    p('The database stores departments, designations, employees, salary structures, attendance summaries, periods, payroll totals, earning and deduction lines, loans, payslips and transfer records. The Python application offers employee CRUD/search, report viewing and CLI commands for attendance, payroll calculation, payslip and transfer record creation. MySQL is the system of record; Python uses parameterized SQL through Connector/Python [6].')
    p('Chapter 2 states the problem, objectives and assumptions. Chapter 3 reviews relational design and database/application documentation. Chapter 4 explains the ER model, normalization, architecture and algorithm. Chapter 5 describes schema, seed data and application modules. Chapter 6 separates expected checks from results actually executed in the available environment. Chapter 7 summarizes learning and future work.')

    h('2. Problem Statement and Objectives')
    h('2.1 Problem definition',2)
    p('The project addresses fragmented payroll information and limited integrity controls in manual or spreadsheet workflows. Employee data may be repeated in each pay period; salary values may be overwritten without preserving the effective period; variable earnings and deductions may be forced into fixed spreadsheet columns; and totals may not reconcile to the underlying line items. A relational system is needed to represent the entities and relationships explicitly, prevent duplicate payroll for the same employee and period, and provide reproducible reports.')
    h('2.2 Objectives',2)
    numbered.counter=0
    for x in ['Design a conceptual ER model covering all required payroll entities and cardinalities.','Convert the model to relations normalized through 3NF and identify key functional dependencies.','Implement MySQL tables with primary, foreign, unique and check constraints, defaults and indexes.','Add realistic fictional data and queries for payroll, deductions, loans, department cost, variance and transfer status.','Build a Python interface connected to MySQL for employee CRUD/search, report display and payroll workflow commands.','Demonstrate transaction handling and specify valid, invalid, duplicate and boundary tests.'] : numbered(x)
    h('2.3 Constraints and assumptions',2)
    p('The database target is MySQL 8.0.16 or later because the design uses enforced CHECK constraints. Python 3.10+ and mysql-connector-python are assumed. All amounts use DECIMAL, not binary floating point. A payroll period is monthly in sample data. Attendance is summarized by period; daily punch records are outside scope. A salary structure is selected by designation and effective dates. The sample formula uses 10% pension on basic, 5% PAYE on gross, and overtime at basic divided by 176 hours. These rates are demonstrative only. Transfer state is manually recorded; no external API is called. Authentication, role-based permissions, tax jurisdiction rules, concurrent payroll locking beyond the loan rows, and production deployment are future work.')
    h('2.4 Functional and quality requirements',2)
    table(['ID','Requirement','Acceptance evidence'],[
      ('FR1','Create, search, update and delete employee records','Tkinter form; foreign keys restrict unsafe deletion'),('FR2','Store salary and attendance inputs','MySQL salary and attendance rows'),('FR3','Calculate payroll and loan repayment','Payroll row plus component rows and updated balance'),('FR4','Record payslip and transfer state','One-to-one linked rows'),('FR5','Run management reports','Views and grouped queries'),('NFR1','Protect data integrity','PK, FK, UNIQUE, CHECK and triggers'),('NFR2','Avoid partial payroll write','Commit/rollback transaction'),('NFR3','Protect query parameters','Parameterized SQL through connector')])

    h('3. Literature Review and Background')
    h('3.1 Relational database foundations',2)
    p('Codd’s relational model [1] describes data using relations and supports operations that combine and filter sets of rows. Payroll demonstrates why this abstraction helps: an employee can be joined to a department, a payroll period and a set of earning lines without copying all descriptive data into every record. Elmasri and Navathe [2] explain conceptual modeling, keys and normalization as core database design practices. Connolly and Begg [3] connect these design ideas to implementation and database lifecycle.')
    h('3.2 Integrity and transaction behavior',2)
    p('The MySQL Reference Manual describes constraints, indexes, views, transactions and triggers used by this implementation [4], [5], [7], [8]. Primary keys distinguish rows; foreign keys preserve references; unique constraints capture candidate keys; and checks validate row-level conditions. Payroll total consistency spans a payroll row and component rows, so the prototype uses triggers for selected conditions and application transaction logic for the complete calculation. InnoDB transactions provide commit and rollback behavior [5].')
    h('3.3 Application integration and safe queries',2)
    p('PEP 249 defines a common Python database API shape, including cursors and transaction conventions [9]. MySQL Connector/Python implements a driver for connecting Python code to MySQL [6]. The Tkinter standard library supports basic desktop forms and tables without introducing a web framework [10]. Decimal arithmetic is used for currency to avoid binary floating-point rounding artifacts [11]. OWASP recommends parameterized queries to separate SQL code from user-supplied values [12].')
    h('3.4 Related-work comparison',2)
    table(['Source','Contribution','Use in this project','Limit'],[
      ('Codd [1]','Relational model','Tables, joins and keys','Foundational, not a payroll implementation'),
      ('Elmasri & Navathe [2]','ER and normalization','Conceptual design and 3NF','General textbook treatment'),
      ('Connolly & Begg [3]','Database systems design','Lifecycle and schema rationale','Not specific to statutory payroll'),
      ('MySQL manual [4]','SQL constraints and schema features','DDL, checks, views and triggers','Behavior depends on server version'),
      ('InnoDB manual [5]','Transactions and foreign keys','Atomic payroll and loan updates','Requires InnoDB engine'),
      ('Connector/Python [6]','Python-MySQL connectivity','DB-API access from the app','Does not supply application rules'),
      ('OWASP [12]','SQL injection prevention','Parameterized statements','Security guidance, not a payroll rules source')])
    p('The reviewed sources provide complementary foundations rather than competing payroll products. Textbooks guide conceptual design, official MySQL documentation describes implementation semantics, and Python/OWASP documentation informs the client connection and query handling. No jurisdiction-specific statutory source is selected because the project brief does not state a country or legal regime.')
    h('3.5 Review of selected sources',2)
    h('3.5.1 Relational data model',3)
    p('Codd’s 1970 paper [1] introduced a relational approach in which data is represented through relations and accessed through operations on sets. The payroll design uses this idea by representing employees, periods and components as tables rather than embedding a complete monthly payroll in a single file. A join assembles related information when a report needs it. This keeps employee and department descriptions from being copied into every payroll record. The paper establishes a model, but it does not prescribe this project’s tables or payroll policy; those are derived from the requirements.')
    h('3.5.2 ER modeling and normalization',3)
    p('Elmasri and Navathe [2] explain conceptual ER modeling, keys and the normal forms that address redundancy and anomalies. This project applies those concepts when it separates department and designation master data, and when it models earning and deduction components as child relations. Connolly and Begg [3] additionally frame database development as a sequence from requirements through design, implementation and maintenance. The three review stages follow that practical lifecycle. Both sources are general database references; they support the design method but do not evaluate the payroll prototype empirically.')
    h('3.5.3 SQL constraints and schema behavior',3)
    p('The MySQL manual [4] describes the concrete syntax and server behavior for constraints. This matters because SQL features vary across database products and versions. The implementation targets MySQL 8.0.16 or later so CHECK constraints are enforced. The manual is the primary source used to check syntax assumptions such as CHECK, foreign key and trigger support. A limitation is that a manual describes product behavior, not whether the modeled business policy is appropriate; local policy must still be confirmed separately.')
    h('3.5.4 Transactions and InnoDB',3)
    p('MySQL’s InnoDB transaction documentation [5] explains transaction boundaries, commit and rollback. Payroll calculation consists of several writes: payroll detail, earning lines, deduction lines and loan updates. Saving only some of these would create a misleading balance. The application therefore commits the writes as a unit and rolls back if an exception occurs. InnoDB documentation supplies the engine behavior, while the application must still choose appropriate transaction boundaries and handle failures correctly.')
    h('3.5.5 Python database connectivity',3)
    p('PEP 249 [9] defines a common Python database API, including connections, cursors and parameter styles. Oracle’s Connector/Python guide [6] documents the specific MySQL connector used here. The two sources help separate general application/database concepts from the driver’s concrete behavior. The prototype uses `%s` placeholders and an explicit commit/rollback sequence. The presence of a driver does not guarantee a correct transaction; that needs to be tested against a running database.')
    h('3.5.6 Desktop interface and arithmetic',3)
    p('The Python Tkinter documentation [10] describes standard widgets for desktop interfaces. The prototype uses forms, buttons and a tree view to avoid introducing a separate web stack. Python’s Decimal documentation [11] supports fixed-point arithmetic suited to currency amounts. Decimal still requires explicit rounding and a documented calculation policy, so the app quantizes percentage and overtime calculations to two decimal places. The interface and arithmetic references are implementation documentation rather than independent performance evaluations.')
    h('3.5.7 SQL safety',3)
    p('OWASP’s SQL Injection Prevention Cheat Sheet [12] recommends prepared statements and parameterized queries. The Tkinter and CLI forms accept employee identifiers and amounts, so interpolating values into SQL text would be unsafe and error-prone. The application binds user values as parameters. This protects the query structure from data values, but it does not replace access control, least-privilege database accounts or input validation; those are future improvements.')

    h('4. Methodology and System Design')
    h('4.1 Development approach',2)
    p('The project follows a small database-first sequence. Requirements identify data facts and user tasks. The ER model captures entity types and relationship cardinalities. Functional dependencies guide normalization. DDL then implements the relations and constraints. Fictional seed records support query demonstrations. Finally, a Python client uses parameterized SQL and a transaction for payroll processing. Each stage has a review document and test evidence checklist.')
    h('4.2 Architecture',2)
    p('USER → Tkinter desktop forms / Python CLI → application validation and payroll calculation → MySQL Connector/Python → MySQL InnoDB database → views and reports',italic=True)
    p('The database is local by default. Environment variables configure the connection. The user interface calls application functions; the application executes parameterized SQL; MySQL enforces relational constraints. The reports query database views and tables. This simple layered arrangement is appropriate for a small student MVP.')
    h('4.3 Entity relationship and cardinality',2)
    pic=doc.add_paragraph();pic.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(str(ROOT/'docs'/'review-1'/'er_diagram.png'),width=Inches(6.2))
    p('Figure 3. Entity relationship diagram showing principal tables, keys and cardinality.')
    p('Department-to-Employee is one-to-many, as is Designation-to-Employee. Designation-to-SalaryStructure is one-to-many over time. Employee-to-Attendance and Employee-to-Loan are one-to-many. Employee and PayrollPeriod form a many-to-many business relation resolved by PayrollDetails. Each payroll can contain multiple earning and deduction rows. A unique payroll_id foreign key in Payslips and BankTransfers makes each association zero-or-one from a payroll. The complete Mermaid ER diagram is included at docs/review-1/er_diagram.md.')
    h('4.4 Normalization and functional dependencies',2)
    p('First normal form is supported by storing scalar values and one earning/deduction component per child row. A fixed design with `earning_1`, `earning_2`, and `earning_3` would impose a changing column set; child rows avoid that repeating group. Second normal form is satisfied because the relations use single-attribute primary keys; the additional candidate key `(employee_id, period_id)` identifies a payroll detail. Third normal form is supported by placing department names and designation descriptions only in their respective tables. Employee rows reference those identifiers instead of carrying repeated descriptive labels. Salary terms are kept in their own effective-dated relation, and period descriptions are separate from payroll results.')
    table(['Relation','Example dependency','Design consequence'],[
      ('departments','department_id → code, name','Department description stored once'),('employees','employee_id → employee_no, name, email, department_id, designation_id','Employee fact references department/designation'),('salary_structures','salary_structure_id → designation, dates, components','Salary history can be effective-dated'),('payroll_periods','period_id → name, dates, status','Period dates stored once'),('payroll_details','payroll_id → employee, period, totals; employee+period → payroll_id','Unique employee-period key'),('payroll_earnings','earning_id → payroll_id, type, amount','Flexible one-row-per-component')])
    h('4.4.1 Anomalies addressed by normalization',3)
    p('If a department name were copied into every employee row, renaming that department would require many updates; a missed row would leave conflicting names. Keeping the name in Departments reduces this update anomaly. If payroll stored repeated columns such as pension_1 and pension_2, adding a deduction type would require a schema change and queries would need to inspect multiple columns. A child table permits a variable number of components and allows a sum over the same amount field. Separating payroll periods prevents each employee payroll from storing a repeated period label and date range. These choices reduce insert, update and delete anomalies while leaving reports able to join the needed descriptions.')
    h('4.5 Payroll algorithm',2)
    numbered.counter=0
    for x in ['Begin a database transaction and retrieve an active employee, matching attendance, payroll period, and the salary structure effective for that period.','Reject missing inputs and an existing payroll row for that employee-period pair.','Compute overtime as basic salary / 176 × overtime hours; calculate gross as basic plus fixed and variable earning lines.','Lock active loan rows, compute each repayment as the lower of balance and installment, and sum repayments.','Compute illustrative pension and PAYE; verify total deductions do not exceed gross; compute net = gross − deductions.','Insert payroll detail and each earning/deduction line; decrement each corresponding loan balance.','Commit the entire transaction. If any SQL or validation error occurs, roll back the payroll rows and loan changes together.'] : numbered(x)
    h('4.5.1 Input and concurrency considerations',3)
    p('Payroll calculation reads several sources that must refer to the same employee and period. A missing attendance row is treated as an input error rather than silently assumed to be full attendance. The selected salary row must cover the payroll period dates. The employee-period UNIQUE constraint is the final protection against duplicates even if two payroll operators start calculations close together. Loan rows are selected FOR UPDATE so a second transaction cannot calculate the same installment against a stale balance before the first transaction commits. If another transaction has already created payroll, the unique constraint raises an error and the calculation rolls back. This is a practical example of using both application checks and database constraints: the initial lookup gives a friendly message, while the database key remains authoritative.')
    h('4.5.2 Currency rounding',3)
    p('All persisted salary and deduction amounts use DECIMAL(12,2). Python uses Decimal values returned by the MySQL connector and rounds calculated overtime, pension and PAYE to two decimal places. The database then stores the same two-decimal scale. Rounding rules can affect one-cent differences, so this prototype uses the Decimal default half-even rounding for quantize operations. A real statutory system would document the legally required rounding rule and test examples at half-cent boundaries. No binary float is used for core payroll arithmetic; input floats only occur in the optional salary maintenance dialog and are sent to MySQL as values, not used in formula calculations.')
    h('4.6 Workflow diagram',2)
    p('Employee + Salary + Attendance + Period → Validate inputs → Calculate earning lines → Calculate deduction/loan lines → Check deduction limit → Insert payroll header and lines → Update loans → COMMIT; on error → ROLLBACK.',italic=True)
    h('4.7 Tools and technologies',2)
    table(['Tool','Purpose','DBMS concept demonstrated'],[
      ('MySQL 8 / InnoDB','Persistent relational storage','DDL, constraints, transactions, views, indexes'),('Python 3','Application logic','Validation and application/database integration'),('Tkinter','Desktop forms and table view','CRUD and search UI'),('mysql-connector-python','Driver','Parameterized SQL and transaction API'),('python-docx / python-pptx','Generate editable report/slides','Submission artifacts')])
    h('4.8 Use cases and data flow',2)
    p('The HR use case is to find an employee, correct contact information and maintain department/designation assignment. The payroll use case begins with a selected period and uses the employee’s effective salary structure and attendance summary. It produces payroll detail and component rows, then updates each loan balance. The finance use case reads deduction and transfer reports and records transfer status. The manager use case reads department cost and variance summaries without changing payroll records.')
    p('Data flow is: user form → field validation → parameterized Python query → MySQL constraints → stored employee or payroll rows → report view → table/report window. During payroll, the calculation service loads source data, locks relevant loan rows, computes totals, writes the payroll and repayment records, and commits. An invalid amount or database error routes to rollback and an error message.')

    h('5. Implementation')
    h('5.1 Database modules and tables',2)
    table(['Table','Main role','Key integrity'],[
      ('departments, designations','Organization master data','PK and unique names/codes'),('employees','Employee master','unique employee_no/email; FK department/designation'),('salary_structures','Effective salary components','FK designation; date and amount checks'),('attendance_summaries','Period attendance','unique employee-period; present day range'),('payroll_periods','Period and status','unique period name/date range'),('loans','Loan principal and balance','FK employee; balance bounded by principal'),('payroll_details','One employee payroll snapshot','unique employee-period; amount checks'),('payroll_earnings/deductions','Variable components','FK payroll; unique type per payroll; nonnegative amount'),('payslips/bank_transfers','Output/payment records','unique payroll FK; transfer amount positive')])
    h('5.1.1 Concise data dictionary',3)
    table(['Table','Representative attributes and meaning'],[
      ('departments','department_id PK; department_code unique; department_name unique; active flag.'),
      ('designations','designation_id PK; designation_name unique; grade; active flag.'),
      ('employees','employee_id PK; employee_no and email unique; names/contact; department_id and designation_id FKs; hire_date/status; bank destination.'),
      ('salary_structures','salary_structure_id PK; designation_id FK; effective_from/to; basic_salary; housing and transport allowances.'),
      ('payroll_periods','period_id PK; unique period_name and date range; payment_date; status.'),
      ('attendance_summaries','attendance_id PK; employee_id FK; period dates; working_days; days_present; overtime_hours.'),
      ('loans','loan_id PK; employee_id FK; type; principal_amount; balance; monthly_installment; start_date/status.'),
      ('payroll_details','payroll_id PK; employee_id, period_id, attendance_id FKs; basic_salary; gross_pay; deductions; net_pay.'),
      ('payroll_earnings','earning_id PK; payroll_id FK; earning_type; amount; unique payroll/type.'),
      ('payroll_deductions','deduction_id PK; payroll_id FK; deduction_type; statutory flag; amount.'),
      ('payslips','payslip_id PK; unique payroll_id FK; generated timestamp; document status.'),
      ('bank_transfers','transfer_id PK; unique payroll_id FK; employee_id FK; bank/account; amount/date/status.')])
    h('5.2 Schema, constraints and indexes',2)
    p('The schema uses InnoDB so foreign keys and transactions are available. AUTO_INCREMENT integer keys make joins simple. Email, employee number and department/designation labels are unique where appropriate. Status fields use enumerated choices. Monetary fields use DECIMAL(12,2). Search indexes cover employee name and department; payroll period and active employee-loan lookups have indexes. The unique `(employee_id, period_id)` index both enforces the rule and supports payroll lookups.')
    p('MySQL CHECK support is enforced starting with MySQL 8.0.16; earlier releases may parse but ignore checks. The setup therefore specifies 8.0.16 or later. Triggers enforce deduction totals against gross and verify net pay when payroll detail is inserted or changed. Since MySQL does not provide deferred constraints across an entire transaction in the same way as some other databases, the application calculates all lines before inserting the payroll header and commits them atomically.')
    h('5.3 Fictional sample dataset',2)
    p('The seed file includes three departments, three designations, three fictional employees, effective salary rows, two periods, attendance summaries, two loans, a January payroll, earning/deduction lines, payslip records and transfer statuses. Names and email addresses use example.test domains; bank account fields contain TEST labels. This prevents accidental representation of actual employee financial data.')
    h('5.4 SQL queries and reports',2)
    p('The reports script demonstrates a view joining payroll details to employees, departments and periods; a grouped deduction sum; a nested query selecting loan balances above an installment benchmark; department gross and net cost; variable-pay variance; and a transfer-status join. It also opens a transaction, changes a status for demonstration and rolls back to prove that the original state remains.')
    h('5.5 Python application modules',2)
    p('The Tkinter window displays employee rows and supports add, select, search, update and delete. A separate salary-structure window supports add, select, update and delete for designation-linked effective salary rows. Required fields and numeric identifiers are checked in the form; MySQL handles uniqueness, references and database-level validation. Parameter placeholders prevent user values from being concatenated into SQL. A foreign-key error during employee delete is shown when dependent payroll or loan records exist.')
    p('The CLI writes attendance summaries, calculates payroll, creates a payslip record and records bank instructions. The payroll service explicitly starts work through a connector transaction, locks active loan rows with `FOR UPDATE`, calculates values with Decimal, writes payroll and component rows, updates the matching loan rows, and commits. Exception handling rolls back and re-raises the error for a visible explanation.')
    h('5.5.1 Employee and salary maintenance',3)
    p('The employee screen presents employee number, name, email, department and designation in a table. Search matches employee number or name. Add and update use bound parameters. The employee number is treated as a stable identifier during update; the UI rejects changing it accidentally. Delete is allowed only when dependent records do not exist, which illustrates referential integrity rather than cascading away historical payroll. A separate salary-structure window lists designation and effective dates with basic, housing and transport values. Add/update forms require numeric values, while database checks ensure positive basic salary and nonnegative allowances. Effective-dated rows preserve prior salary terms instead of overwriting an employee’s historical compensation.')
    h('5.5.2 Attendance and payroll processing',3)
    p('Attendance is captured as one summary per employee and period, with working days, present days and overtime. The unique key prevents duplicate summaries for the same period. The CLI uses an upsert so a correction updates the existing summary rather than creating a duplicate. Payroll calculation requires the period, active employee, matching attendance and effective salary. It calculates salary components and loan deductions in one operation. The detail row contains the basic salary and final totals as a snapshot; separate lines explain how gross and deductions were formed. This makes report queries straightforward and allows a student to inspect a payslip source with normal SELECT statements.')
    h('5.5.3 Loan, payslip and transfer records',3)
    p('Loan repayment for each active loan is capped at its remaining balance. The resulting deduction line is associated with payroll, and the same amount is applied to the corresponding loan row. When a balance reaches zero, a trigger changes the loan status to PAID. Payslip generation creates a status record linked uniquely to a payroll; it does not yet render a PDF. A transfer operation copies the employee’s bank destination and net pay into a uniquely linked bank transfer instruction. Triggers check that transfer employee and amount match payroll. It only records a payment instruction; a finance user must not describe it as a completed bank payment unless an external bank result exists.')
    h('5.5.4 Reporting interface',3)
    p('The desktop interface opens report windows for payroll summary, department cost, loan balances and bank transfers. The SQL behind the screens uses a view or grouped joins rather than reimplementing aggregates in the GUI. This keeps the data query visible and reusable for the review demonstration. Reports are read-only in the interface. The SQL report file remains the source for deduction totals and payroll variance. A manager can compare gross and net cost by department, while finance can filter pending transfer records using a standard WHERE clause.')
    h('5.6 Pseudocode',2)
    p('BEGIN; load employee + salary + attendance + period; if not found then ERROR; if payroll exists then ERROR; compute earnings and gross; lock loans; compute installments; compute statutory deductions; if deductions > gross then ERROR; INSERT payroll; INSERT earnings; INSERT deductions; UPDATE each loan balance; COMMIT. CATCH error: ROLLBACK; report error.')
    h('5.7 Screenshots and outputs',2)
    p('Insert real screenshots here after running the application: (a) employee search and table; (b) add/update form; (c) payroll summary report; (d) MySQL grouped report; (e) duplicate or invalid insert rejected; (f) loan balance before and after payroll. Captions should state the test data and actual result. No screenshots are fabricated in this draft because the workspace has no MySQL server.')
    h('5.8 Validation and error handling',2)
    p('The interface checks required employee fields and parses department/designation identifiers as integers before it sends SQL. Database constraints remain the final authority because users or scripts can bypass the GUI. A duplicate email or employee number returns a connector error that the window displays. Invalid dates, missing parent identifiers, negative values, and attendance day ranges are rejected by SQL data types and constraints. Deleting an employee that still has attendance, loans or payroll is refused by foreign keys, preserving historical references.')
    p('Payroll processing has a transaction boundary around the whole calculation. If the duplicate employee-period check fails, if a deduction exceeds gross, if an inserted component violates a unique key, or if a loan update fails, the exception handler rolls back. Row locks protect the loan rows while the installment is calculated. The unique employee-period key handles a race where two clients attempt to create the same payroll at once: only one insert can succeed, and the other transaction rolls back.')
    h('5.9 Example calculation walkthrough',2)
    p('For the fictional January record for employee EMP001, basic salary is 250,000.00, housing is 50,000.00, transport is 25,000.00, and overtime is 8 hours. The illustrative hourly overtime amount is 250,000 / 176 = 1,420.4545; rounded overtime is 11,363.64. Gross is therefore 336,363.64. The illustrative pension is 10% of basic (25,000.00), and PAYE is 5% of gross (16,818.18). The emergency loan balance at the start of this payroll is 15,000.00 with a 20,000 installment, so repayment is capped at the balance (15,000.00). Total deductions are 56,818.18 and net pay is 279,545.46. The seed data stores the post-period loan balance as zero and the loan status as PAID. These amounts are fictional; the example is included to make arithmetic and database rows reproducible.')

    h('6. Results and Discussion')
    h('6.1 Verification status',2)
    p('The SQL, trigger behavior and application/database flow could not be executed in this workspace because MySQL client/server binaries and a database service are not installed. Therefore this report does not claim successful database execution, measured performance, or completed screenshots. The test matrix lists expected results for the student to run locally. Python source files are checked for syntax separately when building this bundle.')
    h('6.2 Expected functional outputs',2)
    table(['Check','Expected behavior','Evidence needed'],[
      ('Seed and summary view','Three January payroll rows with stored gross, deduction and net totals','MySQL query output'),('Duplicate employee/period','Unique constraint rejects second row','Error screenshot'),('Present days above working days','CHECK rejects row','Error screenshot'),('Net inconsistent with totals','Payroll trigger rejects write','Error screenshot'),('Loan below zero','CHECK/trigger rejects update','Error screenshot'),('Payroll calculation','Net equals gross minus line deductions; loan reduced by installment','Before/after SQL result'),('Report roll-back','Temporary transfer status change disappears after ROLLBACK','Two SELECT outputs')])
    h('6.3 Objective-result mapping',2)
    table(['Objective','Designed implementation','Current evidence status'],[
      ('Model and normalize payroll','ER diagram, separate relations, 3NF analysis','Design documented'),('Protect identity and references','PK/FK/UNIQUE/CHECK constraints','DDL prepared; DB not run here'),('Report payroll and costs','Views, joins, grouping and nested query','SQL prepared; results pending local run'),('Integrate Python with MySQL','Tkinter + connector + payroll service','Source prepared; connection pending local run'),('Use atomic payroll processing','Transaction and rollback path','Code prepared; DB behavior unverified here')])
    h('6.3.1 Interpreting evidence after local execution',3)
    p('When the scripts are run, compare each observed value to its source inputs rather than relying only on a screen total. For the seeded payroll, gross should equal basic salary plus each earning row; deductions should equal the sum of deduction rows; and net should equal gross less deductions. The department query should equal the sum of its employees’ payroll rows. Loan reports should show only positive balances in the active-loan view. Transfer amounts should match payroll net pay. These reconciliations provide simple independent checks using SQL aggregates. If a result differs, inspect the line-item rows first, then the app formula, then any trigger or constraint that affected the transaction.')
    h('6.4 Discussion and limitations',2)
    p('The schema and app illustrate the requested DBMS concepts, but correctness must ultimately be established by execution against the specified MySQL version. Trigger order and connector behavior should be validated on that environment. The payroll policy is intentionally simplified and must be replaced with instructor-approved rules for any real legal jurisdiction. Attendance days are stored but this prototype does not prorate salary for absence. Salary overlaps are not prevented by an exclusion constraint. A bank transfer row does not move money. There is no authentication, encryption-at-rest configuration, audit log, hosted database or web deployment. These limitations are appropriate to disclose in the classroom demo.')
    h('6.5 Test execution record to complete',2)
    p('After MySQL is installed, execute the schema and seed scripts on a disposable database. Record the exact MySQL version, date and command used. For every row in docs/review-2/test_matrix.md, write the observed error or output, mark pass/fail, and attach the screenshot filename. A test is a pass only when the observed behavior matches the expected behavior. If a trigger or CHECK behaves differently from this design, fix and rerun the script before editing this chapter. The present document intentionally does not invent measured outcomes.')
    table(['Test group','Cases to run','Record in final submission'],[
      ('Positive','Seed load; summary view; one payroll calculation; zero overtime boundary','Rows, totals, screenshots'),
      ('Duplicate','Employee number; employee-period payroll','Exact MySQL error and rollback'),
      ('Range','Days present > workdays; negative loan; negative deduction','Constraint/trigger result'),
      ('Arithmetic','Net inconsistent; deductions > gross; transfer amount mismatch','Trigger result and unchanged state'),
      ('Transaction','Force an error after payroll insert but before completion','Show rollback leaves no partial payroll or loan change'),
      ('Application','Add/search/update/delete restriction and reports','Screen capture and interaction notes')])

    h('7. Conclusion and Future Scope')
    h('7.1 Summary and learning outcomes',2)
    p('This project turns payroll requirements into an ER model, normalized MySQL relations, integrity rules, reports and a Python interface. It demonstrates how primary and foreign keys express identities and relationships, how UNIQUE and CHECK constraints reject invalid facts, how views and SQL aggregation support reports, and how transactions keep related payroll and loan changes atomic. The implementation also demonstrates why monetary calculations use decimal arithmetic and why database tests should be supported with visible evidence.')
    p('The student should be able to explain entity selection, cardinality, functional dependencies, 1NF–3NF, candidate keys, indexes, parameterized queries, commit/rollback and the difference between a transfer record and a real bank transfer.')
    h('7.2 Future enhancements',2)
    for x in ['Add jurisdiction-specific statutory rules as effective-dated configuration with faculty-approved source references.','Add role-based authentication and least-privilege database accounts.','Add audit history for salary, employee and payroll changes.','Add salary overlap validation and payroll-period close/reopen workflow.','Generate printable PDF payslips with access control.','Add a web interface only after selecting a secure hosted MySQL service and configuring secrets.'] : bullet(x)

    h('References')
    refs=[
      '[1] E. F. Codd, “A Relational Model of Data for Large Shared Data Banks,” Communications of the ACM, vol. 13, no. 6, pp. 377–387, 1970.',
      '[2] R. Elmasri and S. B. Navathe, Fundamentals of Database Systems, 7th ed. Pearson, 2016.',
      '[3] T. Connolly and C. Begg, Database Systems: A Practical Approach to Design, Implementation, and Management, 6th ed. Pearson, 2014.',
      '[4] Oracle, “MySQL 8.0 Reference Manual: Constraints,” https://dev.mysql.com/doc/refman/8.0/en/constraints.html. Accessed Oct. 16, 2026.',
      '[5] Oracle, “MySQL 8.0 Reference Manual: InnoDB and Transaction Model,” https://dev.mysql.com/doc/refman/8.0/en/innodb-transaction-model.html. Accessed Oct. 16, 2026.',
      '[6] Oracle, “MySQL Connector/Python Developer Guide,” https://dev.mysql.com/doc/connector-python/en/. Accessed Oct. 16, 2026.',
      '[7] Oracle, “MySQL 8.0 Reference Manual: CREATE TRIGGER Statement,” https://dev.mysql.com/doc/refman/8.0/en/create-trigger.html. Accessed Oct. 16, 2026.',
      '[8] Oracle, “MySQL 8.0 Reference Manual: View Processing,” https://dev.mysql.com/doc/refman/8.0/en/view-processing.html. Accessed Oct. 16, 2026.',
      '[9] Python Software Foundation, “PEP 249 — Python Database API Specification v2.0,” https://peps.python.org/pep-0249/. Accessed Oct. 16, 2026.',
      '[10] Python Software Foundation, “ tkinter — Python interface to Tcl/Tk,” https://docs.python.org/3/library/tkinter.html. Accessed Oct. 16, 2026.',
      '[11] Python Software Foundation, “decimal — Decimal fixed point and floating point arithmetic,” https://docs.python.org/3/library/decimal.html. Accessed Oct. 16, 2026.',
      '[12] OWASP Foundation, “SQL Injection Prevention Cheat Sheet,” https://cheatsheetseries.owasp.org/cheatsheets/SQL_Injection_Prevention_Cheat_Sheet.html. Accessed Oct. 16, 2026.',
      '[13] Oracle, “MySQL 8.0 Reference Manual: Optimization and Indexes,” https://dev.mysql.com/doc/refman/8.0/en/optimization-indexes.html. Accessed Oct. 16, 2026.',
      '[14] Oracle, “MySQL 8.0 Reference Manual: CHECK Constraints,” https://dev.mysql.com/doc/refman/8.0/en/create-table-check-constraints.html. Accessed Oct. 16, 2026.',
      '[15] Oracle, “MySQL 8.0 Reference Manual: START TRANSACTION, COMMIT, and ROLLBACK,” https://dev.mysql.com/doc/refman/8.0/en/commit.html. Accessed Oct. 16, 2026.',
      '[16] Oracle, “MySQL 8.0 Reference Manual: Foreign Key Constraints,” https://dev.mysql.com/doc/refman/8.0/en/create-table-foreign-keys.html. Accessed Oct. 16, 2026.',
    ]
    for ref in refs:p(ref)
    h('Appendix A — Setup and Demonstration')
    p('Install MySQL 8.0.16+, load database/schema.sql, database/seed.sql and database/reports.sql in that order, install requirements.txt, configure MYSQL_* environment variables, and run `python -m app.main`. Use `python -m app.cli --help` for available payroll operations. Run the validation script only against a disposable coursework database. See README.md for command examples.')
    p('Submission completion checklist: replace the name, roll and project number placeholders; verify academic year; run every database test; add genuine screenshots and outputs; update the result discussion with measured evidence; check the word/page requirements; and obtain required faculty signatures.')
    doc.save(OUT)

def build_presentation():
    prs=Presentation();prs.slide_width=PptInches(13.333);prs.slide_height=PptInches(7.5)
    navy=RGBColor(20,44,77);teal=RGBColor(20,144,150);pale=RGBColor(238,245,248);white=RGBColor(255,255,255);ink=RGBColor(35,45,58)
    def slide(title,sub=None,bullets=None):
        s=prs.slides.add_slide(prs.slide_layouts[6]);bg=s.background.fill;bg.solid();bg.fore_color.rgb=pale
        bar=s.shapes.add_shape(1,0,0,prs.slide_width,PptInches(.22));bar.fill.solid();bar.fill.fore_color.rgb=teal;bar.line.fill.background()
        tb=s.shapes.add_textbox(PptInches(.65),PptInches(.55),PptInches(12),PptInches(.75)).text_frame;tb.text=title
        for r in tb.paragraphs[0].runs:r.font.name='Aptos Display';r.font.size=PptPt(28);r.font.bold=True;r.font.color.rgb=navy
        if sub:
            st=s.shapes.add_textbox(PptInches(.7),PptInches(1.45),PptInches(11.9),PptInches(.55)).text_frame;st.text=sub
            for r in st.paragraphs[0].runs:r.font.size=PptPt(17);r.font.color.rgb=teal
        if bullets:
            frame=s.shapes.add_textbox(PptInches(.8),PptInches(2.15),PptInches(11.7),PptInches(4.7)).text_frame;frame.word_wrap=True
            for i,b in enumerate(bullets):
                para=frame.paragraphs[0] if i==0 else frame.add_paragraph();para.text=b;para.level=0;para.space_after=PptPt(14)
                for r in para.runs:r.font.name='Aptos';r.font.size=PptPt(21);r.font.color.rgb=ink
        foot=s.shapes.add_textbox(PptInches(.7),PptInches(7.12),PptInches(12),PptInches(.2)).text_frame;foot.text='PBL • MySQL + Python • Siddharth [ROLL NUMBER]'
        for r in foot.paragraphs[0].runs:r.font.size=PptPt(9);r.font.color.rgb=navy
        return s
    slide('Employee Payroll & Statutory Deduction System','Database Management Systems • 25TU03MJM0',['Siddharth [FULL NAME]  |  Roll: [ROLL NUMBER]','Semester 3 • CSE Panthers • Mentor: Dr. Sahithi','Woxsen University • Academic Year 2026–27'])
    slide('Problem and motivation',bullets=['Payroll information spans employees, salary, attendance, earnings, deductions and loans.','Spreadsheet duplication and formula drift make totals harder to audit.','A relational database can connect records and enforce key business rules.'])
    slide('Objectives and scope',bullets=['Maintain employee and effective salary data; record attendance.','Calculate payroll, itemized deductions and loan repayment.','Produce payslip/transfer records and management reports.','Educational prototype only: fictional rates, no live bank connection.'])
    slide('Users and modules',bullets=['HR: employee, department and designation maintenance.','Payroll: salary, attendance, payroll and payslips.','Finance: deductions, loans and transfer status.','Manager: department cost, payroll variance and summary reports.'])
    slide('ER model and key relationships',bullets=['Department 1:N Employee; Designation 1:N Employee and Salary Structure.','Employee 1:N Attendance, Loan and Payroll Detail.','Employee M:N Payroll Period resolved through Payroll Detail.','Payroll Detail 1:N Earnings/Deductions; at most one Payslip and Transfer.','Open docs/review-1/er_diagram.md for the complete ER diagram.'])
    er=prs.slides.add_slide(prs.slide_layouts[6]);er.background.fill.solid();er.background.fill.fore_color.rgb=pale
    bar=er.shapes.add_shape(1,0,0,prs.slide_width,PptInches(.22));bar.fill.solid();bar.fill.fore_color.rgb=teal;bar.line.fill.background()
    titlebox=er.shapes.add_textbox(PptInches(.65),PptInches(.5),PptInches(12),PptInches(.65)).text_frame;titlebox.text='Payroll ER Diagram'
    for run in titlebox.paragraphs[0].runs:run.font.size=PptPt(26);run.font.bold=True;run.font.color.rgb=navy
    er.shapes.add_picture(str(ROOT/'docs'/'review-1'/'er_diagram.png'),PptInches(.45),PptInches(1.35),width=PptInches(12.4))
    slide('Normalization and integrity',bullets=['1NF: scalar facts; each earning/deduction is a separate row.','2NF/3NF: attributes depend on the row key; department/designation descriptions stored once.','PK, FK, UNIQUE, CHECK and DEFAULT protect facts and relationships.','One payroll per employee-period; payroll totals reconciled; loan balance cannot go negative.'])
    slide('MySQL reports and transactions',bullets=['Views and joins: employee gross, deduction and net pay.','GROUP BY: deduction totals and department payroll cost.','Nested query: active loan balances relative to installments.','Variance and bank transfer status reports.','Payroll and loan changes commit together or roll back together.'])
    slide('Python application demo',bullets=['Tkinter employee table: search, add, update and delete.','CLI: attendance, payroll calculation, payslip record and transfer instruction.','Parameterized MySQL Connector/Python queries.','Reports are displayed in the app; MySQL enforces relational constraints.'])
    slide('Validation and honest results',bullets=['Test plan covers duplicate keys, invalid attendance, negative loan and inconsistent payroll totals.','MySQL execution and screenshots must be captured on the student machine; unavailable in this build workspace.','Rates are fictional (10% pension, 5% PAYE, basic/176 overtime); replace only with instructor-approved local policy.'])
    slide('Conclusion and next steps',bullets=['Project demonstrates ER design, 3NF, SQL, constraints, transactions and Python integration.','Run setup and validation scripts; record actual outputs; add genuine screenshots.','Replace personal placeholders and complete faculty certificate/signature.','Future: approved local tax rules, access control, audit history and secure hosted database.'])
    prs.save(PPT)

if __name__=='__main__':
    build_report();build_presentation()
    print(f'Created {OUT} and {PPT}')
